import os
import io
import json
import requests
from docx import Document

BASE_URL = "http://127.0.0.1:8000"

def create_sample_docx():
    doc = Document()
    doc.add_heading("Siddharth Patel - Senior Backend Engineer", 0)
    doc.add_paragraph("Summary: Experienced software engineer specializing in Python, Django, Flask, SQL, REST API, Git, and complex Data Structures.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, Django, Flask, SQL, REST API, Git, Data Structures, Software Development")
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph("Senior Python Developer at TechVentures: Designed and implemented high performance REST APIs and microservices using Python and Flask.")
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.getvalue()

def run_full_suite():
    print("==================================================")
    print("     PHASE 4 COMPLETE SYSTEM VALIDATION SUITE     ")
    print("==================================================")
    
    # 1. Health check
    print("\n1. Testing GET /api/health...")
    r = requests.get(f"{BASE_URL}/api/health")
    assert r.status_code == 200
    h = r.json()
    assert h["status"] == "healthy"
    assert h["model_loaded"] is True
    assert h["database_connected"] is True
    print(f"   [PASS] System healthy! DB Connected: {h['database_connected']}, Model Loaded: {h['model_loaded']}")

    # 2. Recruiter Login
    print("\n2. Testing Recruiter Login (admin@talentai.com)...")
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": "admin@talentai.com", "password": "admin123"})
    assert r.status_code == 200
    token = r.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    print("   [PASS] Recruiter authenticated successfully!")

    # 3. Recruiter Register
    print("\n3. Testing Recruiter Registration...")
    reg_email = f"qa.lead.{abs(hash(os.urandom(6)))}@example.com"
    r = requests.post(f"{BASE_URL}/api/auth/register", json={
        "email": reg_email,
        "password": "qa_password_123",
        "full_name": "QA Lead Engineer",
        "role": "recruiter"
    })
    assert r.status_code == 200
    print(f"   [PASS] New recruiter registered: {reg_email}")

    # 4. Job Descriptions CRUD
    print("\n4. Testing Job Descriptions CRUD...")
    job_payload = {
        "title": "Lead Software Architect (Python/Cloud)",
        "description": "We are seeking a Lead Software Architect with expertise in Python, Django, REST API, SQL, Docker, and Data Structures.",
        "required_skills": ["python", "django", "rest api", "sql", "docker", "data structures"]
    }
    r = requests.post(f"{BASE_URL}/api/jobs", json=job_payload, headers=headers)
    assert r.status_code == 201
    job = r.json()
    jid = job["id"]
    print(f"   Created Job ID {jid}: {job['title']}")
    
    # Update job
    r = requests.put(f"{BASE_URL}/api/jobs/{jid}", json={"title": "Principal Software Architect (Python/Cloud)"}, headers=headers)
    assert r.status_code == 200
    assert r.json()["title"] == "Principal Software Architect (Python/Cloud)"
    print("   [PASS] Job Description CRUD operations verified!")

    # 5. Screening PDF Resume (Real Dataset)
    print("\n5. Testing PDF Resume Screening (from dataset)...")
    pdf_path = os.path.abspath("data/data/INFORMATION-TECHNOLOGY/10089434.pdf")
    with open(pdf_path, "rb") as f:
        files = {"file": ("10089434.pdf", f, "application/pdf")}
        data = {
            "job_description": job["description"],
            "required_skills": json.dumps(job["required_skills"]),
            "job_title": job["title"],
            "job_description_id": jid,
            "candidate_name": "Rohan Verma"
        }
        r = requests.post(f"{BASE_URL}/api/screen-resume", files=files, data=data, headers=headers)
    assert r.status_code == 200
    pdf_res = r.json()
    print(f"   Candidate ID: {pdf_res['candidate_id']} | Category: {pdf_res['predicted_category']} | ATS Score: {pdf_res['ats_score']}%")
    assert pdf_res["ats_score"] >= 0
    assert "screening_id" in pdf_res
    print("   [PASS] PDF Resume screened & stored in SQLite!")

    # 6. Screening DOCX Resume
    print("\n6. Testing DOCX Resume Screening...")
    docx_bytes = create_sample_docx()
    files = {"file": ("Siddharth_Patel_Resume.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {
        "job_description": job["description"],
        "required_skills": json.dumps(["python", "django", "flask", "sql", "rest api", "git", "data structures", "software development"]),
        "job_title": job["title"],
        "candidate_name": "Siddharth Patel"
    }
    r = requests.post(f"{BASE_URL}/api/screen-resume", files=files, data=data, headers=headers)
    assert r.status_code == 200
    docx_res = r.json()
    print(f"   Candidate: Siddharth Patel | Category: {docx_res['predicted_category']} | ATS Score: {docx_res['ats_score']}% | Rec: {docx_res['recommendation']}")
    assert docx_res["ats_score"] >= 70.0  # High match
    print("   [PASS] DOCX Resume screened with Strong Match score!")

    # 7. Screening TXT Resume
    print("\n7. Testing TXT Resume Screening...")
    txt_content = (
        "ANANYA IYER - Python Web Developer\n"
        "Expert in Python, Flask, Django, SQL databases, Git version control, REST API engineering and Data Structures.\n"
        "Over 5 years building scalable distributed software."
    )
    files = {"file": ("Ananya_Iyer_Resume.txt", io.BytesIO(txt_content.encode("utf-8")), "text/plain")}
    data = {
        "job_description": job["description"],
        "required_skills": json.dumps(["python", "flask", "django", "sql", "git", "rest api", "data structures"]),
        "candidate_name": "Ananya Iyer"
    }
    r = requests.post(f"{BASE_URL}/api/screen-resume", files=files, data=data, headers=headers)
    assert r.status_code == 200
    txt_res = r.json()
    print(f"   Candidate: Ananya Iyer | Category: {txt_res['predicted_category']} | ATS Score: {txt_res['ats_score']}% | Rec: {txt_res['recommendation']}")
    print("   [PASS] TXT Resume screened successfully!")

    # 8. Screening History Verification
    print("\n8. Testing Screening History (GET /api/screenings)...")
    r = requests.get(f"{BASE_URL}/api/screenings", headers=headers)
    assert r.status_code == 200
    history = r.json()
    print(f"   Total History Records in DB: {history['total']}")
    assert history["total"] >= 3
    print("   [PASS] Screening History audit trail verified!")

    # 9. Resumes List Verification
    print("\n9. Testing Resumes List (GET /api/resumes)...")
    r = requests.get(f"{BASE_URL}/api/resumes", headers=headers)
    assert r.status_code == 200
    resumes = r.json()
    assert resumes["total"] >= 3
    print(f"   Total Resumes Recorded: {resumes['total']}")
    print("   [PASS] Resumes file management verified!")

    # 10. Candidate Details with History
    print("\n10. Testing Candidate Profile (GET /api/candidates/{id})...")
    cid = docx_res["candidate_id"]
    r = requests.get(f"{BASE_URL}/api/candidates/{cid}")
    assert r.status_code == 200
    cdata = r.json()
    print(f"   Candidate: {cdata['name']} (ID: {cdata['id']}) | Category: {cdata['category']} | Matched: {len(cdata['matched_skills'])} skills")
    assert len(cdata["matched_skills"]) > 0
    print("   [PASS] Candidate Details profile verified!")

    # 11. CSV Export Endpoints
    print("\n11. Testing CSV Export Endpoints...")
    r1 = requests.get(f"{BASE_URL}/api/export/rankings")
    r2 = requests.get(f"{BASE_URL}/api/export/shortlist")
    r3 = requests.get(f"{BASE_URL}/api/export/screenings")
    assert r1.status_code == 200 and "text/csv" in r1.headers.get("content-type", "")
    assert r2.status_code == 200 and "text/csv" in r2.headers.get("content-type", "")
    assert r3.status_code == 200 and "text/csv" in r3.headers.get("content-type", "")
    print(f"   [PASS] CSV exports generated! Rankings: {len(r1.text)} bytes, Shortlist: {len(r2.text)} bytes, History: {len(r3.text)} bytes")

    # 12. Error handling checks
    print("\n12. Testing Defensive Error Handling...")
    # Empty file
    r = requests.post(f"{BASE_URL}/api/screen-resume", files={"file": ("empty.txt", io.BytesIO(b""), "text/plain")}, data={"job_description": "test"})
    assert r.status_code == 400
    # Empty JD
    r = requests.post(f"{BASE_URL}/api/screen-resume", files={"file": ("doc.txt", io.BytesIO(b"sample resume text here with python"), "text/plain")}, data={"job_description": ""})
    print("   Empty JD status:", r.status_code, r.text[:80])
    assert r.status_code in [400, 422]
    # Invalid extension
    r = requests.post(f"{BASE_URL}/api/screen-resume", files={"file": ("malicious.exe", io.BytesIO(b"binary"), "application/octet-stream")}, data={"job_description": "test role description"})
    print("   Invalid ext status:", r.status_code, r.text[:80])
    assert r.status_code in [400, 422]
    print("   [PASS] All error conditions rejected with safe HTTP status codes and friendly messages!")

    print("\n==================================================")
    print("     ALL 12 TEST SUITE CATEGORIES PASSED (100%)    ")
    print("==================================================")

if __name__ == "__main__":
    run_full_suite()
