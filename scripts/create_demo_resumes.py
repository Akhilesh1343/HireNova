import os
from docx import Document

out_dir = os.path.abspath("data/demo_resumes")
os.makedirs(out_dir, exist_ok=True)

# 1. Alex Rivera DOCX Resume
doc = Document()
doc.add_heading("Alex Rivera", 0)
p = doc.add_paragraph("Full Stack Developer | Python, SQL & Web Technologies")
p.bold = True
doc.add_paragraph("Email: alex.rivera@example.com | Portfolio: alexrivera.dev | Location: Austin, TX")

doc.add_heading("Professional Summary", level=1)
doc.add_paragraph(
    "Versatile software development professional with 4 years of experience building modern web applications. "
    "Skilled in Python, SQL databases, Git version control, and frontend frameworks. Passionate about writing clean, "
    "maintainable software and collaborating with engineering teams."
)

doc.add_heading("Technical Skills", level=1)
doc.add_paragraph("Programming: Python, JavaScript, TypeScript, SQL")
doc.add_paragraph("Web Frameworks: Flask, Express.js, React, REST API")
doc.add_paragraph("Databases: PostgreSQL, SQLite")
doc.add_paragraph("Developer Tools: Git, GitHub, Docker, Postman")
doc.add_paragraph("Core Competencies: Software Development, Problem Solving, Data Structures")

doc.add_heading("Work Experience", level=1)
doc.add_paragraph("Software Developer | Apex Solutions (2021 - Present)")
doc.add_paragraph("- Developed and maintained web applications utilizing Python and Flask backends with React frontend.")
doc.add_paragraph("- Structured complex SQL queries and migrations for relational databases.")
doc.add_paragraph("- Utilized Git for daily version control, pull requests, and continuous deployment workflows.")
doc.add_paragraph("- Contributed to ongoing software development initiatives and technical documentation.")

doc.add_paragraph("Junior Programmer | CodeCraft Labs (2020 - 2021)")
doc.add_paragraph("- Wrote automated test scripts using Python.")
doc.add_paragraph("- Assisted in building SQL reports and database backups.")
doc.add_paragraph("- Participated in weekly code reviews and team standups.")

doc.add_heading("Education", level=1)
doc.add_paragraph("B.S. in Information Technology, State University (2016 - 2020)")

docx_path = os.path.join(out_dir, "Alex_Rivera_FullStack.docx")
doc.save(docx_path)
print(f"Created: {docx_path}")

# 2. Jordan Hayes Marketing TXT Resume
txt_content = """JORDAN HAYES
Digital Marketing & Creative Content Specialist
Email: jordan.hayes@example.com | Phone: +1 (555) 789-0123

PROFESSIONAL SUMMARY
Creative marketing manager with 5+ years driving audience engagement, social media campaigns, and brand storytelling. Expert in SEO, Google Ads, content marketing, and Adobe Creative Cloud. Seeking to leverage analytical and strategic creative skills.

CORE COMPETENCIES
- Social Media Strategy, Paid Advertising (Google, Meta, LinkedIn)
- SEO Optimization, Copywriting, Public Relations
- Analytics: Google Analytics 4, Tableau, Excel
- Creative: Adobe Photoshop, Illustrator, Premiere Pro

EXPERIENCE
Marketing Manager | GrowthWave Agency (2021 - Present)
- Led end-to-end creative marketing campaigns generating 300% ROI.
- Managed quarterly ad spend across multiple digital channels.
- Spearheaded team branding and social media content creation.

Marketing Coordinator | BrightStar Media (2019 - 2021)
- Authored weekly blog posts and marketing emails.
- Monitored engagement metrics and prepared client presentations.

EDUCATION
B.A. in Communications & Media Studies (2015 - 2019)
"""

txt_path = os.path.join(out_dir, "Jordan_Hayes_Marketing.txt")
with open(txt_path, "w", encoding="utf-8") as f:
    f.write(txt_content)
print(f"Created: {txt_path}")

# 3. Create a README in data/demo_resumes explaining datasets
readme_content = """# Demo & Test Resumes

This directory contains standalone, clearly-labelled demo candidate resumes designed for live testing and college demonstration:

1. `Elena_Rostova_Python_Expert.txt` — Strong match candidate with deep Python, Django, Flask, SQL, REST API, Git, Data Structures, and Software Development experience.
2. `Alex_Rivera_FullStack.docx` — Good match full-stack engineer with Python, SQL, REST API, Git, and web technologies.
3. `Jordan_Hayes_Marketing.txt` — Unmatched / Low match resume (Digital Marketing domain) to demonstrate category classification and low ATS scoring.

---
### Dataset Distinctions:
- **Original Dataset**: Located in `data/data/` (24 categorized folders with 2,484 resumes) and `Resume/Resume.csv`. Must remain untouched.
- **Demo / Test Data**: Located in `data/demo_resumes/` for manual drag-and-drop or file upload testing.
- **User-Uploaded Resumes**: Temporarily processed by the backend and tracked in the SQLite database (`backend/screening.db`).
"""

readme_path = os.path.join(out_dir, "README.md")
with open(readme_path, "w", encoding="utf-8") as f:
    f.write(readme_content)
print(f"Created: {readme_path}")
